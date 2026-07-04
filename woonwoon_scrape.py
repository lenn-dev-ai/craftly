#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""woonwoon_scrape.py — Scraper für woonwoon.de (Wohnungs-Inserate) → Excel.

woonwoon.de ist ein WordPress-Portal (immomakler-Theme) für möblierte
Wohnungen auf Zeit in Berlin. Dieses Skript:

  1. crawlt die Inserats-Übersicht /apartments/ (URL-Pagination),
  2. liest pro Inserat *alle* verfügbaren Felder aus
     (Detailzeilen `data-*`, Ausstattung, Beschreibung, Geo-Koordinaten,
     Bilder, Preis/Kaution/Verfügbarkeit …),
  3. schreibt das Ergebnis in eine .xlsx (eine Zeile je Inserat).

Die Spalten sind dynamisch: jedes je gesehene Detailfeld wird zu einer
Spalte (Reihenfolge = erstes Auftreten). So landet wirklich alles im Sheet,
ohne dass das Feld-Set vorab festgelegt werden muss.

Nutzung:
    python3 woonwoon_scrape.py                 # alles -> woonwoon.xlsx
    python3 woonwoon_scrape.py --out foo.xlsx  # eigenes Zielfile
    python3 woonwoon_scrape.py --limit 5       # nur 5 Inserate (Test)
    python3 woonwoon_scrape.py --delay 1.5     # langsamer/höflicher

Abhängigkeiten:  pip install requests openpyxl
"""

from __future__ import annotations

import argparse
import html as ihtml
import re
import sys
import time
from collections import OrderedDict

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

BASE = "https://woonwoon.de"
ARCHIVE = BASE + "/apartments/"
HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/124.0 Safari/537.36"
    ),
    "Accept-Language": "de-DE,de;q=0.9,en;q=0.7",
}

# Detail-Links wie /apartments/wohnung-in-berlin-mieten-20e22a/
LISTING_RE = re.compile(r'href="(https://woonwoon\.de/apartments/[a-z0-9-]+/)"')
# Pseudo-Links, die keine Inserate sind
SKIP_SLUGS = {"feed", "merkliste"}

# Feste Spalten am Anfang (Reihenfolge bleibt erhalten)
FIXED_COLS = ["url", "Titel", "Stadt/Bezirk", "Lat", "Lng", "Ausstattung",
              "Bilder (Anzahl)", "Titelbild", "Bild-URLs", "Beschreibung"]


def clean(s: str) -> str:
    """HTML-Tags entfernen, Entities auflösen, Whitespace normalisieren."""
    s = re.sub(r"<[^>]+>", " ", s)
    s = ihtml.unescape(s)
    return re.sub(r"\s+", " ", s).strip()


def fetch(session: requests.Session, url: str, tries: int = 3) -> str | None:
    """GET mit kleinen Retries + Backoff. Gibt HTML-Text oder None zurück."""
    for attempt in range(1, tries + 1):
        try:
            r = session.get(url, headers=HEADERS, timeout=30)
            if r.status_code == 404:
                return None
            r.raise_for_status()
            return r.text
        except requests.RequestException as exc:
            if attempt == tries:
                print(f"  ! Fehler bei {url}: {exc}", file=sys.stderr)
                return None
            time.sleep(1.5 * attempt)
    return None


def collect_listing_urls(session: requests.Session, delay: float,
                         max_pages: int = 200) -> list[str]:
    """Alle Inserats-URLs über die paginierte Übersicht sammeln.

    Stoppt, sobald eine Seite keine *neuen* URLs mehr liefert (woonwoon
    wiederholt jenseits der letzten Seite die letzten Treffer).
    """
    urls: "OrderedDict[str, None]" = OrderedDict()
    page = 1
    empty_streak = 0
    while page <= max_pages:
        page_url = ARCHIVE if page == 1 else f"{ARCHIVE}page/{page}/"
        html = fetch(session, page_url)
        if html is None:
            break
        found = []
        for m in LISTING_RE.finditer(html):
            u = m.group(1)
            slug = u.rstrip("/").rsplit("/", 1)[-1]
            if slug in SKIP_SLUGS:
                continue
            found.append(u)
        new = [u for u in found if u not in urls]
        for u in new:
            urls[u] = None
        print(f"  Seite {page}: {len(found)} Links, {len(new)} neu "
              f"(gesamt {len(urls)})")
        # Keine neuen Treffer mehr -> Ende (2x zur Sicherheit)
        if not new:
            empty_streak += 1
            if empty_streak >= 2:
                break
        else:
            empty_streak = 0
        page += 1
        time.sleep(delay)
    return list(urls.keys())


def parse_detail(html: str, url: str) -> "OrderedDict[str, str]":
    """Ein Inserat in ein Feld->Wert-Dict zerlegen (alle verfügbaren Daten)."""
    rec: "OrderedDict[str, str]" = OrderedDict()
    rec["url"] = url

    # Titel
    m = re.search(r'class="property-title[^"]*"[^>]*>(.*?)</', html, re.S)
    rec["Titel"] = clean(m.group(1)) if m else ""

    # Detailzeilen: <li class="list-group-item data-KEY"> <div dt>Label</div>
    #                                                     <div dd>Wert</div> </li>
    for li in re.findall(
        r'<li class="list-group-item[^"]*"[^>]*>(.*?)</li>', html, re.S
    ):
        dt = re.search(r'class="dt[^"]*"[^>]*>(.*?)</div>', li, re.S)
        dd = re.search(r'class="dd[^"]*"[^>]*>(.*?)</div>', li, re.S)
        if not dt or not dd:
            continue
        label = clean(dt.group(1))
        value = clean(dd.group(1))
        if label and value:
            rec[label] = value

    # Ausstattung (property-features: ✓ Balkon …)
    feats = []
    fm = re.search(r'property-features panel.*?(<ul.*?</ul>)', html, re.S)
    if fm:
        for li in re.findall(r'<li[^>]*>(.*?)</li>', fm.group(1), re.S):
            c = clean(li).lstrip("✓ ").strip()
            if c:
                feats.append(c)
    rec["Ausstattung"] = ", ".join(feats)

    # Geo-Koordinaten (im eingebetteten Karten-JSON)
    lat = re.search(r'latitude"\s*:\s*"?([-0-9.]+)', html)
    lng = re.search(r'longitude"\s*:\s*"?([-0-9.]+)', html)
    rec["Lat"] = lat.group(1) if lat else ""
    rec["Lng"] = lng.group(1) if lng else ""

    # Stadt/Bezirk aus Adresse ableiten (z. B. "10961 Berlin (Kreuzberg)")
    addr = rec.get("Adresse", "")
    city = re.search(r"\d{4,5}\s+(.*)$", addr)
    rec["Stadt/Bezirk"] = clean(city.group(1)) if city else ""

    # Bilder: Vollbilder (Thumbnails mit -NNNxNNN-Suffix verwerfen)
    all_imgs = re.findall(
        r"https://woonwoon\.de/wp-content/uploads/[^\s\"'<>]+?\.(?:jpg|jpeg|png|webp)",
        html,
    )
    full = []
    for u in all_imgs:
        if re.search(r"-\d+x\d+\.(?:jpg|jpeg|png|webp)$", u):
            continue  # skalierte Variante
        if u not in full:
            full.append(u)
    rec["Bilder (Anzahl)"] = str(len(full))
    rec["Bild-URLs"] = " | ".join(full)
    og = re.search(r'<meta property="og:image" content="([^"]+)"', html)
    rec["Titelbild"] = og.group(1) if og else (full[0] if full else "")

    # Beschreibung
    dm = re.search(
        r'property-description panel.*?panel-body[^>]*>(.*?)</div>', html, re.S
    )
    if dm:
        desc = clean(dm.group(1))
        desc = re.sub(r"^Beschreibung\s*", "", desc)
        rec["Beschreibung"] = desc
    else:
        rec["Beschreibung"] = ""

    return rec


def write_xlsx(records: list[dict], out_path: str) -> None:
    """Datensätze mit dynamischen Spalten als .xlsx schreiben."""
    # Spalten-Reihenfolge: feste Spalten zuerst, dann alle übrigen
    # Felder in der Reihenfolge ihres ersten Auftretens.
    columns: list[str] = list(FIXED_COLS)
    seen = set(columns)
    for rec in records:
        for k in rec:
            if k not in seen:
                seen.add(k)
                columns.append(k)

    wb = Workbook()
    ws = wb.active
    ws.title = "woonwoon"

    header_font = Font(bold=True)
    ws.append(columns)
    for cell in ws[1]:
        cell.font = header_font
        cell.alignment = Alignment(vertical="top")

    for rec in records:
        ws.append([rec.get(col, "") for col in columns])

    # Kopfzeile fixieren + Autofilter
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}1"

    # Spaltenbreiten grob an Inhalt anpassen (gedeckelt)
    for i, col in enumerate(columns, start=1):
        width = len(str(col))
        for rec in records:
            width = max(width, len(str(rec.get(col, ""))))
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 60)

    wb.save(out_path)


def main() -> int:
    ap = argparse.ArgumentParser(description="Scrapt woonwoon.de-Inserate nach Excel.")
    ap.add_argument("--out", default="woonwoon.xlsx", help="Ziel-.xlsx (Default: woonwoon.xlsx)")
    ap.add_argument("--delay", type=float, default=0.8, help="Pause (Sek.) zwischen Requests")
    ap.add_argument("--limit", type=int, default=0, help="Nur N Inserate (0 = alle, für Tests)")
    ap.add_argument("--max-pages", type=int, default=200, help="Sicherheits-Limit Übersichtsseiten")
    args = ap.parse_args()

    session = requests.Session()

    print("» Sammle Inserats-URLs …")
    urls = collect_listing_urls(session, args.delay, args.max_pages)
    if args.limit:
        urls = urls[: args.limit]
    print(f"» {len(urls)} Inserate gefunden.\n")

    records = []
    for idx, url in enumerate(urls, start=1):
        print(f"[{idx}/{len(urls)}] {url}")
        html = fetch(session, url)
        if html is None:
            print("  ! übersprungen (kein Inhalt)")
            continue
        try:
            rec = parse_detail(html, url)
            records.append(rec)
            print(f"    -> {rec.get('Titel','(ohne Titel)')[:70]}")
        except Exception as exc:  # einzelnes Inserat darf den Lauf nicht killen
            print(f"  ! Parse-Fehler: {exc}", file=sys.stderr)
        time.sleep(args.delay)

    if not records:
        print("Keine Datensätze — nichts geschrieben.", file=sys.stderr)
        return 1

    write_xlsx(records, args.out)
    print(f"\n✓ {len(records)} Inserate -> {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
